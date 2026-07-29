"""Convert EDGR Markdown manuscripts to DOCX (Times New Roman, tables, math as italic)."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

ROOT = Path(__file__).resolve().parent
FIG = ROOT / "figures"

BOLD_ITALIC_CODE = re.compile(
    r"(\*\*[^*]+\*\*|\*[^*\n]+\*|`[^`]+`)"
)


def set_run_font(run, name: str = "Times New Roman", size: int = 12, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_formatted(paragraph, text: str, size: int = 12) -> None:
    parts = BOLD_ITALIC_CODE.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            r = paragraph.add_run(part[2:-2])
            set_run_font(r, size=size, bold=True)
        elif part.startswith("*") and part.endswith("*") and len(part) >= 3 and not part.startswith("**"):
            r = paragraph.add_run(part[1:-1])
            set_run_font(r, size=size, italic=True)
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            set_run_font(r, name="Consolas", size=max(9, size - 1))
        else:
            r = paragraph.add_run(part)
            set_run_font(r, size=size)


def is_table_sep(line: str) -> bool:
    return bool(re.match(r"^\|?\s*:?-{3,}", line.strip()))


def parse_table(lines: list[str], i: int):
    rows = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        line = lines[i].strip()
        if is_table_sep(line):
            i += 1
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        rows.append(cells)
        i += 1
    return rows, i


def try_add_figure(doc: Document, svg_path: Path, width_in: float = 6.2) -> bool:
    png_path = svg_path.with_suffix(".png")
    if not png_path.exists():
        try:
            import cairosvg

            cairosvg.svg2png(url=str(svg_path), write_to=str(png_path), output_width=1400)
        except Exception:
            return False
    if png_path.exists():
        doc.add_picture(str(png_path), width=Inches(width_in))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        return True
    return False


def md_to_docx(md_path: Path, out_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    i = 0
    in_code = False
    code_buf: list[str] = []
    para_buf: list[str] = []

    def flush_para() -> None:
        nonlocal para_buf
        if not para_buf:
            return
        raw = " ".join(para_buf).strip()
        para_buf = []
        if not raw or raw == "---":
            return
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.first_line_indent = Cm(0.75)
        add_formatted(p, raw, size=12)
        m = re.search(r"\*File:\*\s*`([^`]+)`", raw)
        if m:
            name = Path(m.group(1)).name
            svg = FIG / name
            if svg.exists():
                if not try_add_figure(doc, svg):
                    note = doc.add_paragraph()
                    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r = note.add_run(
                        f"[Figure: {name} — open SVG in figures/ and paste into Word if needed]"
                    )
                    set_run_font(r, size=10, italic=True)

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.5)
                r = p.add_run("\n".join(code_buf))
                set_run_font(r, name="Consolas", size=9)
                code_buf = []
                in_code = False
            else:
                flush_para()
                in_code = True
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if line.strip() == "---":
            flush_para()
            i += 1
            continue

        if line.strip().startswith("\\["):
            flush_para()
            math_lines: list[str] = []
            s = line.strip()
            if s == "\\[":
                i += 1
                while i < len(lines) and lines[i].strip() != "\\]":
                    math_lines.append(lines[i].strip())
                    i += 1
                i += 1
            elif s.endswith("\\]"):
                math_lines = [s[2:-2].strip()]
                i += 1
            else:
                math_lines = [s[2:]]
                i += 1
                while i < len(lines) and not lines[i].strip().endswith("\\]"):
                    math_lines.append(lines[i].strip())
                    i += 1
                if i < len(lines):
                    math_lines.append(lines[i].strip()[:-2].strip())
                    i += 1
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(" ".join(math_lines))
            set_run_font(r, size=11, italic=True)
            continue

        if line.strip().startswith("|") and i + 1 < len(lines) and is_table_sep(lines[i + 1].strip()):
            flush_para()
            rows, i = parse_table(lines, i)
            if rows:
                cols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=cols)
                table.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci in range(cols):
                        cell = table.rows[ri].cells[ci]
                        cell.text = ""
                        p = cell.paragraphs[0]
                        val = row[ci] if ci < len(row) else ""
                        add_formatted(p, val, size=10)
                        if ri == 0:
                            for run in p.runs:
                                run.bold = True
                doc.add_paragraph()
            continue

        hm = re.match(r"^(#{1,6})\s+(.*)$", line)
        if hm:
            flush_para()
            level = len(hm.group(1))
            title = hm.group(2).strip()
            if level == 1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(12)
                add_formatted(p, title, size=16)
                for run in p.runs:
                    run.bold = True
            else:
                p = doc.add_heading(level=min(level - 1, 3))
                p.clear()
                add_formatted(p, title, size=14 if level == 2 else 12)
                for run in p.runs:
                    run.bold = True
            i += 1
            continue

        if re.match(r"^[-*]\s+", line) or re.match(r"^\d+\.\s+", line):
            flush_para()
            content = re.sub(r"^[-*]\s+", "", line)
            content = re.sub(r"^\d+\.\s+", "", content)
            p = doc.add_paragraph(style="List Bullet")
            add_formatted(p, content.strip(), size=12)
            i += 1
            continue

        if not line.strip():
            flush_para()
            i += 1
            continue

        para_buf.append(line.strip())
        i += 1

    flush_para()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"Wrote {out_path} ({out_path.stat().st_size:,} bytes)")


def main() -> None:
    pairs = [
        (
            ROOT / "EDGR_international_paper_MDPI_format.md",
            ROOT / "EDGR_international_paper_SUBMIT.docx",
        ),
        (
            ROOT / "EDGR_bai_bao_dinh_dang_PDF_mau.md",
            ROOT / "EDGR_bai_bao_tieng_Viet_SUBMIT.docx",
        ),
    ]
    for src, dst in pairs:
        if not src.exists():
            print("Missing", src)
            continue
        md_to_docx(src, dst)


if __name__ == "__main__":
    main()
